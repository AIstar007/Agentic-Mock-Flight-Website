from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/home/user/project'); DOCS=ROOT/'docs'; OUT=Path('/home/user/emirates_multi_agent_solution_design.docx')
MAROON='8B152B'; DARK='2C2024'; RED='C21D3A'; GOLD='C6A35A'; GREY='665E5D'; WHITE='FFFFFF'
def font(size,bold=False):
    options=['/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    for f in options:
        if Path(f).exists(): return ImageFont.truetype(f,size)
    return ImageFont.load_default()
def create_cover_art(path):
    w,h=1450,760; im=Image.new('RGB',(w,h)); px=im.load()
    for y in range(h):
        for x in range(w):
            t=(x/w)*.42+(1-y/h)*.12; px[x,y]=(int(48+46*t),int(16+10*t),int(27+19*t))
    d=ImageDraw.Draw(im); d.rectangle((0,585,w,h),fill=(42,19,27))
    buildings=[(25,505,110,625),(130,470,210,625),(230,525,290,625),(300,440,390,625),(405,500,478,625),(495,455,580,625),(600,520,675,625),(695,490,760,625),(780,415,855,625),(875,480,960,625),(980,500,1050,625),(1070,455,1150,625),(1170,510,1250,625),(1270,475,1350,625),(1370,525,1440,625)]
    for i,box in enumerate(buildings): d.rectangle(box,fill=(63+i%3*5,31+i%4*3,38+i%3*4))
    d.polygon([(1000,625),(1038,360),(1055,240),(1064,360),(1100,625)],fill=(95,57,55)); d.line((1055,240,1055,175),fill=(218,174,106),width=4)
    plane=[(180,310),(580,245),(760,125),(812,133),(710,230),(1070,176),(1160,195),(1075,230),(735,290),(580,430),(520,430),(612,310),(270,370)]
    d.polygon(plane,fill=(237,233,228)); d.polygon([(180,310),(735,290),(580,430),(520,430),(612,310),(270,370)],fill=(194,29,58)); d.polygon([(610,300),(1073,229),(1125,208),(1160,195),(1075,230),(735,290)],fill=(208,166,92)); d.line([(185,315),(576,252),(750,132)],fill=(255,255,255),width=5)
    d.text((80,55),'EMIRATES',font=font(68,True),fill=(255,255,255)); d.text((85,133),'PASSENGER JOURNEY · MULTI-AGENT POC',font=font(27,True),fill=(224,190,132)); d.text((85,681),'SOLUTION DESIGN  /  TECHNICAL ARCHITECTURE  /  VERSION 1.0',font=font(22,True),fill=(242,225,217)); im.save(path)
def create_architecture_art(path):
    w,h=1500,850; im=Image.new('RGB',(w,h),(255,255,255)); d=ImageDraw.Draw(im)
    d.text((38,22),'EMIRATES PASSENGER JOURNEY POC · CURRENT LOGICAL ARCHITECTURE',font=font(28,True),fill=(139,21,43)); d.text((40,64),'Synthetic workflow only · arrows show logical calls, not live airline integration',font=font(17),fill=(95,82,82))
    cols=[('1. CHANNEL',[('Passenger','Browser / mobile'),('Booking UI','Search + PDP'),('Check-in UI','PNR + documents'),('Floating chat','Typed + Sarvam voice')]),('2. API GATEWAY',[('FastAPI','REST endpoints'),('Pydantic','Typed schemas'),('Session auth','Token + state'),('Country API','ISO name / codes')]),('3. ORCHESTRATION',[('Workflow router','Task + page context'),('Slot parser','Current pending field'),('Action contract','Allow-listed actions'),('State manager','Preserve / advance')]),('4. SPECIALISTS',[('Search / booking','Mock commerce'),('Passenger rules','DOB + companion'),('Seats / extras','Per pax / leg'),('Eligibility','Ordered gates')]),('5. ADAPTERS',[('Mock Navitaire','Flights + seat map'),('Mock Timatic','Synthetic document rules'),('FAQ / RAG','Evidence-aware answers'),('Sarvam proxy','Saaras + Bulbul REST')]),('6. RESULTS',[('Country display','Full name + codes'),('Document result','PASS / FAIL / REVIEW'),('Review card','Reason + solution'),('Boarding passes','Synthetic only')])]
    x0=30; gap=14; cw=(w-60-gap*5)//6; ytop=125
    for ci,(head,boxes) in enumerate(cols):
        x=x0+ci*(cw+gap); d.rounded_rectangle((x,ytop,x+cw,ytop+54),radius=11,fill=(139,21,43)); d.text((x+10,ytop+14),head,font=font(17,True),fill=(255,255,255)); yy=ytop+68
        for j,(title,sub) in enumerate(boxes):
            fill=(250,244,240) if j%2==0 else (246,241,237); d.rounded_rectangle((x,yy,x+cw,yy+105),radius=8,fill=fill,outline=(211,194,183),width=2); d.text((x+10,yy+16),title,font=font(16,True),fill=(54,40,41))
            words=sub.split(); lines=[]; cur=''
            for word in words:
                if len(cur+' '+word)>22: lines.append(cur); cur=word
                else: cur=(cur+' '+word).strip()
            if cur: lines.append(cur)
            for k,line in enumerate(lines[:3]): d.text((x+10,yy+48+k*21),line,font=font(13),fill=(101,85,85))
            yy+=117
        if ci<5:
            ax=x+cw+1; ay=ytop+300; bx=ax+gap-2; d.line((ax,ay,bx,ay),fill=(198,163,90),width=4); d.polygon([(bx,ay),(bx-8,ay-6),(bx-8,ay+6)],fill=(198,163,90))
    y=650; d.rounded_rectangle((38,y,w-38,y+140),radius=14,fill=(46,33,38)); d.text((62,y+18),'CROSS-CUTTING CONTROLS',font=font(19,True),fill=(224,190,132))
    labels=['Input validation + deterministic gates','Session scope + audit metadata','No secrets in browser / logs','Mock-only status + safe abstention','Current-question voice cancellation']
    for i,label in enumerate(labels):
        x=62+(i%3)*475; yy=y+58+(i//3)*35; d.ellipse((x,yy+3,x+11,yy+14),fill=(194,29,58)); d.text((x+22,yy),label,font=font(15),fill=(250,242,238))
    im.save(path)
DOCS.joinpath('assets').mkdir(parents=True,exist_ok=True); cover=DOCS/'assets'/'cover-illustration.png'; arch=DOCS/'assets'/'architecture.png'; create_cover_art(cover); create_architecture_art(arch)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def set_cell_margins(cell,top=75,start=80,bottom=75,end=80):
    tcPr=cell._tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for edge,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+edge))
        if node is None: node=OxmlElement('w:'+edge); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
def set_bottom_border(paragraph,color=GOLD,size='8'):
    pPr=paragraph._p.get_or_add_pPr(); borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),size); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),color); borders.append(bottom); pPr.append(borders)
def add_field(run,field):
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=f' {field} '; sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); txt=OxmlElement('w:t'); txt.text='1'; end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for el in (begin,instr,sep,txt,end): run._r.append(el)
def set_run(run,size=None,bold=False,color=DARK,italic=False,name='Aptos'):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    if size: run.font.size=Pt(size)
    run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color); return run
def add_rich_text(p,text,size=8.5,color=DARK):
    pattern=re.compile(r'(\*\*.*?\*\*|`[^`]+`)'); pos=0
    for m in pattern.finditer(text):
        if m.start()>pos: set_run(p.add_run(text[pos:m.start()]),size=size,color=color)
        token=m.group()
        if token.startswith('**'): set_run(p.add_run(token[2:-2]),size=size,bold=True,color=color)
        else: set_run(p.add_run(token[1:-1]),size=size,color=MAROON,name='Consolas')
        pos=m.end()
    if pos<len(text): set_run(p.add_run(text[pos:]),size=size,color=color)
def fmt_p(p,before=0,after=2.5,line=1.0,keep=False):
    f=p.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line
    if keep: f.keep_with_next=True
def add_body(doc,text,size=8.3,after=2.5):
    p=doc.add_paragraph(); fmt_p(p,after=after,line=1.03); add_rich_text(p,text,size=size); return p
def add_heading(doc,text,level=2):
    p=doc.add_paragraph(); fmt_p(p,before=5 if level==2 else 3,after=2,line=1.0,keep=True)
    size,color={1:(16.5,MAROON),2:(10.5,RED),3:(9.1,MAROON)}[level]; set_run(p.add_run(text),size=size,bold=True,color=color)
    if level==1: set_bottom_border(p,GOLD,'8')
    return p
def add_bullet(doc,text,numbered=False):
    p=doc.add_paragraph(style='List Number' if numbered else 'List Bullet'); fmt_p(p,after=1.2,line=1.0); p.paragraph_format.left_indent=Inches(.21); p.paragraph_format.first_line_indent=Inches(-.12); add_rich_text(p,text,size=8.1); return p
def add_table(doc,rows):
    if not rows: return
    cols=max(len(r) for r in rows); table=doc.add_table(rows=0,cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths={2:[2.05,4.55],3:[1.4,2.05,3.15],4:[1.1,1.6,1.8,2.1]}.get(cols,[6.6/cols]*cols)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for j,cell in enumerate(cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
            value=row[j] if j<len(row) else ''; p=cell.paragraphs[0]; fmt_p(p,after=0,line=1.0); add_rich_text(p,value,size=7.2 if cols>=3 else 7.7,color=WHITE if ri==0 else DARK)
            if ri==0: shade(cell,MAROON)
            elif ri%2==0: shade(cell,'FBF6F2')
        if ri==0: set_repeat_table_header(table.rows[0])
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=Pt(1)
def create_header_footer(sec):
    sec.header_distance=Inches(.28); sec.footer_distance=Inches(.26)
    p=sec.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; fmt_p(p,after=1,line=1.0); set_run(p.add_run('EMIRATES'),size=10,bold=True,color=MAROON); set_run(p.add_run('   |   PASSENGER JOURNEY POC'),size=7.7,bold=True,color=GREY); set_bottom_border(p,GOLD,'7')
    p=sec.footer.paragraphs[0]; fmt_p(p,after=0,line=1.0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(p.add_run('SOLUTION DESIGN  •  SYNTHETIC DATA ONLY   |   PAGE '),size=7,color=GREY); add_field(p.add_run(),'PAGE'); set_run(p.add_run(' OF '),size=7,color=GREY); add_field(p.add_run(),'NUMPAGES')
def render_page(doc,raw,page_num):
    lines=raw.strip().splitlines(); i=0
    if any(x.strip()=='[[COVER]]' for x in lines):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,after=0); set_run(p.add_run('EMIRATES'),size=30,bold=True,color=MAROON)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,after=6); set_run(p.add_run('PASSENGER JOURNEY  /  MULTI-AGENT POC'),size=10,bold=True,color=GOLD)
        doc.add_picture(str(cover),width=Inches(6.65)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.paragraphs[-1].paragraph_format.space_after=Pt(5)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,before=5,after=4,line=1.0); set_run(p.add_run('Multi-Agent Flight Booking & Web Check-in POC'),size=22,bold=True,color=MAROON)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,after=7); set_run(p.add_run('Solution Design and Technical Architecture'),size=14.3,bold=True,color=DARK)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,after=4); set_run(p.add_run('Version 1.0   •   12 October 2026'),size=9.8,bold=True,color=RED)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,after=7); set_run(p.add_run('Prepared by Alen Thomas  |  Internal demonstrator'),size=8.7,color=GREY)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt_p(p,before=3,after=0); set_run(p.add_run('NOT AN OFFICIAL AIRLINE SERVICE · SYNTHETIC DATA ONLY'),size=7.7,bold=True,color=MAROON)
        return
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip(): i+=1; continue
        if line.startswith('# '):
            add_heading(doc,line[2:].strip(),1); i+=1
            if page_num==17:
                doc.add_picture(str(arch),width=Inches(6.75)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.paragraphs[-1].paragraph_format.space_after=Pt(4)
        elif line.startswith('## '): add_heading(doc,line[3:].strip(),2); i+=1
        elif line.startswith('### '): add_heading(doc,line[4:].strip(),3); i+=1
        elif line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?',c.replace(' ','')) for c in cells): block.append(cells)
                i+=1
            add_table(doc,block)
        elif line.startswith('- '):
            while i<len(lines) and lines[i].strip().startswith('- '): add_bullet(doc,lines[i].strip()[2:]); i+=1
        elif re.match(r'^\d+\.\s+',line):
            while i<len(lines) and re.match(r'^\d+\.\s+',lines[i].strip()): add_bullet(doc,re.sub(r'^\d+\.\s+','',lines[i].strip()),True); i+=1
        else:
            chunk=[line]; i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('# ','## ','### ','|','- ')) and not re.match(r'^\d+\.\s+',lines[i].strip()):
                chunk.append(lines[i].strip()); i+=1
            add_body(doc,' '.join(chunk),size=8.15 if page_num==3 else 8.35,after=2.5)

def main():
    raw=(DOCS/'solution_design.md').read_text(encoding='utf-8')
    pages=[part.strip() for part in raw.split('===PAGE===') if part.strip()]
    if len(pages)!=27: raise SystemExit(f'Expected 27 content pages; found {len(pages)}')
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=Inches(.64); sec.bottom_margin=Inches(.54); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    create_header_footer(sec)
    normal=doc.styles['Normal']; normal.font.name='Aptos'; normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Aptos'); normal.font.size=Pt(8.35); normal.font.color.rgb=RGBColor.from_string(DARK)
    for name in ['List Bullet','List Number']:
        st=doc.styles[name]; st.font.name='Aptos'; st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Aptos'); st.font.size=Pt(8.1)
    doc.core_properties.title='Emirates Passenger Journey POC — Solution Design and Technical Architecture'; doc.core_properties.subject='Multi-agent flight booking and web check-in demonstrator'; doc.core_properties.author='Alen Thomas'; doc.core_properties.keywords='Emirates-inspired POC; multi-agent; booking; web check-in; Sarvam; ISO-3166; mock document validation'
    for idx,page in enumerate(pages,start=1):
        if idx>1: doc.add_page_break()
        render_page(doc,page,idx)
    doc.save(OUT)
    print(f'CREATED={OUT}\nCONTENT_BLOCKS={len(pages)}\nBYTES={OUT.stat().st_size}')
if __name__=='__main__': main()
